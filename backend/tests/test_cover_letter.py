import asyncio
from pathlib import Path

import pytest
from docx import Document

from app import llm
from app.services.cover_letter import CoverLetterError, write_cover_letter, write_cover_letter_docx
from app.services.cover_letter.generate import LetterOutput, style_problems
from app.services.cv.language import find_letter_banned

GOOD = [
    "At Acme Health Trust I work as a Data Scientist, building demand forecasting models in Python with XGBoost "
    "and checking them with time-based validation. That is close to the model building your advert describes. "
    "The part I like most is the moment a rough question turns into something a team can actually use.",
    "Most of my recent work runs on Azure. Models are deployed with MLflow, and I keep an eye on how they perform "
    "over time, which lines up with the monitoring you ask for. A good share of my time goes on explaining results "
    "to clinical and operational stakeholders and turning what we find into changes to the service.",
    "Before moving into data science I spent many years in healthcare, and I hold an MSc in Data Science from the "
    "University of Sussex. Healthcare taught me to work with people who are busy and sceptical, and who want a "
    "plain answer rather than a long one. I haven't worked on fraud detection, so I would be learning that side of "
    "the role. The modelling and the communication, though, are things I do every day.",
    "Thank you for reading this. I would be glad to talk about the role in more detail whenever suits you. "
    "A short call would be fine, or a longer conversation if that is easier for your team.",
]


def run(coro):
    return asyncio.run(coro)


def patch_llm(monkeypatch, *outputs):
    calls = {"n": 0}

    async def fake(system, user, schema, max_tokens=0):
        out = outputs[min(calls["n"], len(outputs) - 1)]
        calls["n"] += 1
        return LetterOutput(paragraphs=list(out))

    monkeypatch.setattr(llm, "complete_json", fake)
    return calls


def make(sample_cv, sample_job, sample_score, tmp_path):
    return run(write_cover_letter(sample_cv, sample_job, sample_score, tmp_path))


def test_clean_letter_is_saved_with_txt_and_docx(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    calls = patch_llm(monkeypatch, GOOD)
    letter = make(sample_cv, sample_job, sample_score, tmp_path)
    assert calls["n"] == 1
    assert letter.text.startswith("Dear Hiring Manager,")
    assert letter.text.rstrip().endswith("Kind regards,\nJANE EXAMPLE")
    assert (tmp_path / "cover_letter.txt").read_text(encoding="utf-8") == letter.text
    doc = Document(letter.docx_path)
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "JANE EXAMPLE" in body and "jane@example.com" in body and "07000000000" in body
    assert "Kind regards," in body


def test_invented_employer_is_blocked(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    bad = [*GOOD[:-1], "Earlier I spent two years at Google leading the search quality team.", GOOD[-1]]
    patch_llm(monkeypatch, bad)
    letter = make(sample_cv, sample_job, sample_score, tmp_path)
    assert "Google" not in letter.text


def test_invented_metric_is_blocked(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    bad = [*GOOD, "My models cut forecasting error by 40% across the trust."]
    patch_llm(monkeypatch, bad)
    letter = make(sample_cv, sample_job, sample_score, tmp_path)
    assert "40%" not in letter.text


def test_real_number_from_cv_is_allowed(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    good = [*GOOD, "One piece of work saved around £1M a year, and I would like to do more like it."]
    patch_llm(monkeypatch, good)
    assert "£1M" in make(sample_cv, sample_job, sample_score, tmp_path).text


def test_claiming_missing_skill_is_blocked(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    bad = [*GOOD, "I have used Kubernetes extensively to run my models in production."]
    patch_llm(monkeypatch, bad)
    letter = make(sample_cv, sample_job, sample_score, tmp_path)
    assert "Kubernetes" not in letter.text


def test_honest_gap_is_kept(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    patch_llm(monkeypatch, GOOD)
    assert "I haven't worked on fraud detection" in make(sample_cv, sample_job, sample_score, tmp_path).text


def test_invented_company_fact_is_blocked(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    bad = [*GOOD, "I admire Orbit Bank for its Sustainability Programme and its award-winning culture."]
    patch_llm(monkeypatch, bad)
    letter = make(sample_cv, sample_job, sample_score, tmp_path)
    assert "Sustainability" not in letter.text


def test_seniority_and_years_claims_are_blocked(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    bad = [*GOOD, "As a principal data scientist with 10 years of machine learning experience I lead teams."]
    patch_llm(monkeypatch, bad)
    letter = make(sample_cv, sample_job, sample_score, tmp_path)
    assert "principal" not in letter.text.lower()
    assert "10 years" not in letter.text


def test_retry_gets_feedback_and_clean_second_draft_is_used(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    bad = [*GOOD, "I spent three years at Google."]
    seen = []

    async def fake(system, user, schema, max_tokens=0):
        seen.append(user)
        return LetterOutput(paragraphs=bad if len(seen) == 1 else GOOD)

    monkeypatch.setattr(llm, "complete_json", fake)
    letter = make(sample_cv, sample_job, sample_score, tmp_path)
    assert len(seen) == 2 and "Google" in seen[1] and "problems" in seen[1]
    assert "Google" not in letter.text


def test_llm_unavailable_gives_clear_error_and_no_files(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    async def boom(*a, **k):
        raise RuntimeError("usage limit reached")

    monkeypatch.setattr(llm, "complete_json", boom)
    with pytest.raises(CoverLetterError, match="unavailable"):
        make(sample_cv, sample_job, sample_score, tmp_path)
    assert not list(Path(tmp_path).glob("cover_letter.*"))


def test_nothing_usable_left_raises(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    junk = [
        "I led a team of 30 engineers at Google.",
        "I won an award at Microsoft in 2019.",
    ]
    patch_llm(monkeypatch, junk)
    with pytest.raises(CoverLetterError, match="No cover letter"):
        make(sample_cv, sample_job, sample_score, tmp_path)
    assert not list(Path(tmp_path).glob("cover_letter.*"))


def test_cliche_text_is_flagged_and_removed(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    cliche = "I am a passionate, results-driven team player with a proven track record of leveraging data."
    assert len(find_letter_banned(cliche)) >= 4
    patch_llm(monkeypatch, [*GOOD, cliche])
    letter = make(sample_cv, sample_job, sample_score, tmp_path)
    for phrase in ("passionate", "results-driven", "team player", "proven track record"):
        assert phrase not in letter.text.lower()


def test_first_person_is_not_flagged():
    assert find_letter_banned("I worked on forecasting models.") == []


def test_style_problems_detected():
    same = ["I do this. I do that. I do more.", "I did this. I did that.", "I will do this. I will go."]
    found = " ".join(style_problems(same))
    assert "same word" in found and 'start with "I"' in found
    assert "em-dash" in " ".join(style_problems(["one — two — three " * 40]))


def test_docx_regenerates_from_edited_text(tmp_path, sample_cv):
    out = write_cover_letter_docx("Dear Hiring Manager,\n\nEdited by the user.\n\nKind regards,\nJANE EXAMPLE", sample_cv, tmp_path / "x.docx")
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "Edited by the user." in text and "jane@example.com" in text
