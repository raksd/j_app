import asyncio

from docx import Document

from app import llm
from app.services.cv import parse_cv, tailor_cv, write_cv_docx
from app.services.cv.guard import TailorBullet, TailorOutput, TailorRole, check_bullet, check_summary, guard_skills, cv_text
from app.services.cv.language import find_banned, shared_openers


def run(coro):
    return asyncio.run(coro)


def identity(cv) -> TailorOutput:
    return TailorOutput(
        summary=cv.summary,
        skills=list(cv.skills),
        experience=[
            TailorRole(index=i, bullets=[TailorBullet(source=j, text=b) for j, b in enumerate(r.bullets)])
            for i, r in enumerate(cv.experience)
        ],
    )


def patch(monkeypatch, *outputs):
    calls = {"n": 0}

    async def fake(system, user, schema, max_tokens=0):
        out = outputs[min(calls["n"], len(outputs) - 1)]
        calls["n"] += 1
        return out

    monkeypatch.setattr(llm, "complete_json", fake)
    return calls


def tailor(cv, job, score, tmp_path):
    return run(tailor_cv(cv, job, score, tmp_path))


# ---- parsing / writing -------------------------------------------------------------

def test_write_then_parse_round_trip(tmp_path, sample_cv):
    path = write_cv_docx(sample_cv, tmp_path / "cv.docx")
    back = parse_cv(path)
    assert back.full_name == sample_cv.full_name
    assert back.email == sample_cv.email and back.phone == sample_cv.phone
    assert back.skills == sample_cv.skills
    assert [r.company for r in back.experience] == [r.company for r in sample_cv.experience]
    assert back.experience[0].bullets == sample_cv.experience[0].bullets
    assert back.experience[1].bullets == sample_cv.experience[1].bullets
    assert back.summary == sample_cv.summary


def test_parse_reads_cv_laid_out_in_a_table(tmp_path):
    doc = Document()
    doc.add_paragraph("JANE EXAMPLE")
    doc.add_paragraph("jane@example.com | 07000000000")
    doc.add_paragraph("Career profile")
    doc.add_paragraph("Data scientist.")
    doc.add_paragraph("Professional experience")
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.paragraphs[0].text = "Acme Health Trust | Feb 2023 – Present"
    cell.add_paragraph("Data Scientist")
    cell.add_paragraph("Built forecasting models in Python.")
    doc.add_paragraph("Education")
    doc.add_paragraph("University of Sussex — MSc Data Science")
    path = tmp_path / "t.docx"
    doc.save(str(path))
    cv = parse_cv(path)
    assert cv.email == "jane@example.com"
    assert cv.experience and "Built forecasting models in Python." in " ".join(cv.experience[0].bullets)


# ---- honesty guard -----------------------------------------------------------------

def test_guard_blocks_invented_metric_skill_job_cert_and_years(sample_cv):
    whole = cv_text(sample_cv)
    src = sample_cv.experience[0].bullets[0]
    assert check_bullet("Build forecasting models in Python and improved accuracy by 40%.", src, whole)
    assert check_bullet("Build forecasting models in Python and Kubernetes.", src, whole)
    assert check_bullet("Build forecasting models in Python while at Google.", src, whole)
    assert check_summary("Data Scientist who is AWS Certified in machine learning.", sample_cv)
    assert check_summary("Data Scientist with 10 years of experience in machine learning.", sample_cv)
    assert check_summary("Principal Data Scientist working in machine learning.", sample_cv)


def test_guard_accepts_faithful_rewrite(sample_cv):
    whole = cv_text(sample_cv)
    src = sample_cv.experience[0].bullets[2]
    assert check_bullet("Explain findings to clinical and operational stakeholders and turn them into service changes.", src, whole) is None
    assert check_summary("Data Scientist with around 4 years of experience in machine learning and predictive modelling.", sample_cv) is None


def test_buzzwords_flagged_in_new_text_but_not_when_already_in_source(sample_cv):
    assert "results-driven" in find_banned("A results-driven analyst")
    whole = cv_text(sample_cv)
    assert check_bullet("Leveraged Python to build models.", "Built models in Python.", whole)
    assert check_bullet("Leveraged Python to build models.", "Leveraged Python to build models.", whole) is None


def test_skills_guard_removes_invented_and_restores_dropped(sample_cv):
    proposed = [
        "Programming & Data: Python, Rust, SQL",
        "Machine Learning: XGBoost",
        "Invented Heading: Spark",
    ]
    result, notes = guard_skills(proposed, sample_cv)
    joined = " | ".join(result)
    assert "Rust" not in joined and "Spark" not in joined and "Invented Heading" not in joined
    for line in sample_cv.skills:
        for item in line.split(":", 1)[1].split(","):
            assert item.strip() in joined
    assert notes


def test_shared_openers():
    assert shared_openers(["Built a", "Built b", "Built c", "Built d"])
    assert not shared_openers(["Built a", "Led b", "Ran c", "Wrote d"])


# ---- tailor_cv end to end with a mocked LLM -----------------------------------------

def test_faithful_tailoring_is_applied(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    out = identity(sample_cv)
    out.experience[0].bullets.reverse()
    out.skills = list(reversed(sample_cv.skills))
    patch(monkeypatch, out)
    result = tailor(sample_cv, sample_job, sample_score, tmp_path)
    assert result.cv.skills[0].startswith("Cloud & Engineering")
    assert result.cv.experience[0].bullets[0] == sample_cv.experience[0].bullets[3]
    assert sorted(result.cv.experience[0].bullets) == sorted(sample_cv.experience[0].bullets)
    assert (tmp_path / "cv.docx").exists()


def test_invented_content_never_reaches_the_cv(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    out = identity(sample_cv)
    out.summary = "Data Scientist with 10 years of experience, AWS Certified, leading teams at Google."
    out.skills = [*sample_cv.skills[:2], "Programming & Data: Python, SQL, Pandas, Kubernetes, Rust", sample_cv.skills[2]]
    out.experience[0].bullets[0].text = "Build forecasting models and improved accuracy by 40% at Google."
    patch(monkeypatch, out)
    result = tailor(sample_cv, sample_job, sample_score, tmp_path)
    text = cv_text(result.cv)
    for bad in ("10 years", "AWS", "Google", "40%", "Kubernetes", "Rust"):
        assert bad not in text
    assert result.cv.summary == sample_cv.summary
    assert result.cv.experience[0].bullets[0] == sample_cv.experience[0].bullets[0]
    assert result.cv.education == sample_cv.education
    assert result.cv.certifications == sample_cv.certifications
    assert [(r.title, r.company, r.start, r.end) for r in result.cv.experience] == [
        (r.title, r.company, r.start, r.end) for r in sample_cv.experience
    ]


def test_new_job_in_output_is_ignored(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    out = identity(sample_cv)
    out.experience.append(TailorRole(index=5, bullets=[TailorBullet(source=0, text="Led a team at Google.")]))
    patch(monkeypatch, out)
    result = tailor(sample_cv, sample_job, sample_score, tmp_path)
    assert len(result.cv.experience) == 2
    assert "Google" not in cv_text(result.cv)


def test_mostly_rejected_rewrite_returns_original_with_note(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    out = identity(sample_cv)
    for b in out.experience[0].bullets:
        b.text += " It raised accuracy by 99% at Google."
    patch(monkeypatch, out)
    result = tailor(sample_cv, sample_job, sample_score, tmp_path)
    assert result.cv == sample_cv
    assert "unchanged" in result.changes[0].lower()
    assert (tmp_path / "cv.docx").exists()


def test_llm_failure_returns_original_cv(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    async def boom(*a, **k):
        raise RuntimeError("usage limit")

    monkeypatch.setattr(llm, "complete_json", boom)
    result = tailor(sample_cv, sample_job, sample_score, tmp_path)
    assert result.cv == sample_cv
    assert "unavailable" in result.changes[0]
    assert (tmp_path / "cv.docx").exists()


def test_repair_retry_receives_problem_feedback(monkeypatch, tmp_path, sample_cv, sample_job, sample_score):
    bad = identity(sample_cv)
    bad.experience[0].bullets[0].text = "Build forecasting models and improved accuracy by 40%."
    users = []

    async def fake(system, user, schema, max_tokens=0):
        users.append(user)
        return bad if len(users) == 1 else identity(sample_cv)

    monkeypatch.setattr(llm, "complete_json", fake)
    result = tailor(sample_cv, sample_job, sample_score, tmp_path)
    assert len(users) == 2 and "40%" in users[1]
    assert result.cv.experience[0].bullets[0] == sample_cv.experience[0].bullets[0]
