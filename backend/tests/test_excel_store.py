from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from app.models import ApplicationRecord, ApplicationStatus
from app.storage import excel_store


def make_record(id_="r1", **kw):
    data = dict(
        id=id_,
        date_applied=date(2026, 1, 5),
        company="Acme",
        job_title="Engineer",
        location="London",
        work_mode="remote",
        job_url="https://example.com/job/1",
        ats_score=72.5,
        requirements_not_matched=["Kubernetes", "Go"],
        cv_file="output/j1/cv.docx",
        cover_letter_file="output/j1/cover_letter.txt",
        last_updated=datetime(2026, 1, 5, 10, 0, 0),
    )
    data.update(kw)
    return ApplicationRecord(**data)


@pytest.fixture
def path(tmp_path):
    return tmp_path / "tracker.xlsx"


def test_creates_file_and_round_trips(path):
    assert excel_store.list_records(path) == []
    assert path.exists()
    rec = make_record()
    excel_store.append_record(rec, path)
    excel_store.append_record(make_record("r2", company="Beta"), path)
    records = excel_store.list_records(path)
    assert [r.id for r in records] == ["r1", "r2"]
    assert records[0] == rec
    assert records[1].company == "Beta"


def test_layout_and_validation(path):
    excel_store.append_record(make_record(), path)
    ws = load_workbook(path)[excel_store.SHEET]
    assert [c.value for c in ws[1]] == excel_store.HEADERS
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    assert ws["G2"].hyperlink.target == "https://example.com/job/1"
    assert ws["I2"].value == "Kubernetes; Go"
    validations = ws.data_validations.dataValidation
    assert len(validations) == 1
    dv = validations[0]
    assert dv.type == "list"
    assert dv.showErrorMessage
    assert dv.errorStyle == "stop"
    assert "Invalid status" in (dv.errorTitle or "")
    for status in ApplicationStatus:
        assert status.value in dv.formula1
    assert "L2" in str(dv.sqref)


def test_update_status(path):
    excel_store.append_record(make_record(), path)
    updated = excel_store.update_status("r1", ApplicationStatus.INTERVIEW, "Phone screen", path)
    assert updated.status == ApplicationStatus.INTERVIEW
    assert updated.notes == "Phone screen"
    stored = excel_store.list_records(path)[0]
    assert stored.status == ApplicationStatus.INTERVIEW
    assert stored.last_updated >= datetime(2026, 1, 5, 10, 0, 0)
    excel_store.update_status("r1", ApplicationStatus.OFFER, None, path)
    assert excel_store.list_records(path)[0].notes == "Phone screen"


def test_update_unknown_id(path):
    excel_store.append_record(make_record(), path)
    with pytest.raises(excel_store.RecordNotFoundError):
        excel_store.update_status("nope", ApplicationStatus.OFFER, None, path)


def test_atomic_write_leaves_original_on_failure(path, monkeypatch):
    excel_store.append_record(make_record(), path)
    before = path.read_bytes()

    def boom(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr(excel_store.os, "replace", boom)
    with pytest.raises(excel_store.TrackerLockedError, match="open in Excel"):
        excel_store.append_record(make_record("r2"), path)
    assert path.read_bytes() == before
    assert [p.name for p in path.parent.iterdir()] == ["tracker.xlsx"]


def test_list_records_tolerates_manual_status_and_mode_values(path):
    excel_store.append_record(make_record(), path)
    wb = load_workbook(path)
    ws = wb[excel_store.SHEET]
    ws.cell(row=2, column=12).value = "Phone screen"
    ws.cell(row=2, column=6).value = "WFH"
    wb.save(path)
    wb.close()
    records = excel_store.list_records(path)
    assert len(records) == 1
    assert records[0].status == ApplicationStatus.APPLIED
    assert records[0].work_mode == "unknown"


def set_cell(path, column, value):
    wb = load_workbook(path)
    wb[excel_store.SHEET].cell(row=2, column=column).value = value
    wb.save(path)
    wb.close()


@pytest.mark.parametrize(
    "column,value,field,expected",
    [
        (2, "5 Jan 2026", "date_applied", date(2026, 1, 5)),
        (2, "05/01/2026", "date_applied", date(2026, 1, 5)),
        (2, 46027, "date_applied", date(2026, 1, 5)),
        (8, "72.5%", "ats_score", 72.5),
        (8, "N/A", "ats_score", 0.0),
        (8, None, "ats_score", 0.0),
        (14, "yesterday", "last_updated", datetime(2026, 1, 5, 0, 0)),
    ],
)
def test_hand_typed_cells_are_repaired_not_fatal(path, column, value, field, expected):
    excel_store.append_record(make_record(), path)
    set_cell(path, column, value)
    records = excel_store.list_records(path)
    assert len(records) == 1
    assert getattr(records[0], field) == expected


def test_unreadable_row_is_skipped_and_others_still_load(path):
    excel_store.append_record(make_record("r1"), path)
    excel_store.append_record(make_record("r2", company="Beta"), path)
    wb = load_workbook(path)
    ws = wb[excel_store.SHEET]
    ws.cell(row=2, column=2).value = "not a date"
    ws.cell(row=2, column=14).value = "not a date either"
    wb.save(path)
    wb.close()
    records = excel_store.list_records(path)
    assert [r.id for r in records] == ["r2"]


def test_blank_date_falls_back_to_last_updated(path):
    excel_store.append_record(make_record(), path)
    set_cell(path, 2, None)
    records = excel_store.list_records(path)
    assert records[0].date_applied == date(2026, 1, 5)


def test_existing_tracker_without_validation_gains_it_on_next_write(path):
    excel_store.append_record(make_record(), path)
    wb = load_workbook(path)
    ws = wb[excel_store.SHEET]
    ws.data_validations.dataValidation = []
    wb.save(path)
    wb.close()
    assert not load_workbook(path)[excel_store.SHEET].data_validations.dataValidation

    excel_store.update_status("r1", ApplicationStatus.INTERVIEW, None, path)

    validations = load_workbook(path)[excel_store.SHEET].data_validations.dataValidation
    assert len(validations) == 1
    assert validations[0].showErrorMessage
    assert validations[0].errorStyle == "stop"


def test_repeated_writes_do_not_stack_validations(path):
    excel_store.append_record(make_record(), path)
    excel_store.append_record(make_record("r2"), path)
    excel_store.update_status("r1", ApplicationStatus.OFFER, None, path)
    assert len(load_workbook(path)[excel_store.SHEET].data_validations.dataValidation) == 1
