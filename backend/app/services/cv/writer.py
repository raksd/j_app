from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.models import StructuredCV

FONT = "Calibri"


def _heading(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(3)
    run = para.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    para._p.get_or_add_pPr().append(borders)


def _bullet(doc: Document, text: str) -> None:
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(2)
    para.add_run(text)


def _dates(start: str, end: str) -> str:
    return f"{start} – {end}" if start and end else (start or end)


def write_cv_docx(cv: StructuredCV, out_path: Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.75)
        section.top_margin = section.bottom_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    name = doc.add_paragraph()
    name.paragraph_format.space_after = Pt(0)
    run = name.add_run(cv.full_name)
    run.bold = True
    run.font.size = Pt(18)
    if cv.headline:
        doc.add_paragraph(cv.headline).paragraph_format.space_after = Pt(0)
    contact = [p for p in (cv.location, cv.phone, cv.email, *cv.links) if p]
    if contact:
        doc.add_paragraph(" | ".join(contact))

    if cv.summary:
        _heading(doc, "Career profile")
        doc.add_paragraph(cv.summary)

    if cv.skills:
        _heading(doc, "Core competencies")
        for line in cv.skills:
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(1)
            if ":" in line:
                label, items = line.split(":", 1)
                para.add_run(label + ":").bold = True
                para.add_run(items)
            else:
                para.add_run(line)

    if cv.experience:
        _heading(doc, "Professional experience")
        for role in cv.experience:
            head = doc.add_paragraph()
            head.paragraph_format.space_before = Pt(4)
            head.paragraph_format.space_after = Pt(0)
            company = ", ".join(p for p in (role.company, role.location) if p)
            head.add_run(company).bold = True
            dates = _dates(role.start, role.end)
            if dates:
                head.add_run("  |  " + dates)
            if role.title:
                title = doc.add_paragraph()
                title.paragraph_format.space_after = Pt(2)
                title.add_run(role.title).italic = True
            for bullet in role.bullets:
                _bullet(doc, bullet)

    if cv.projects:
        _heading(doc, "Selected projects")
        for project in cv.projects:
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(2)
            para.add_run(project.name + (":" if project.description else "")).bold = True
            if project.description:
                para.add_run(" " + project.description)
            for bullet in project.bullets:
                _bullet(doc, bullet)

    if cv.education:
        _heading(doc, "Education & qualifications")
        for edu in cv.education:
            text = " — ".join(p for p in (edu.institution, edu.qualification) if p)
            dates = _dates(edu.start, edu.end)
            if dates:
                text += f" ({dates})"
            if edu.details:
                text += f", {edu.details}"
            doc.add_paragraph(text).paragraph_format.space_after = Pt(1)

    if cv.certifications:
        _heading(doc, "Certifications")
        for cert in cv.certifications:
            doc.add_paragraph(cert).paragraph_format.space_after = Pt(1)

    doc.save(str(out_path))
    return out_path
