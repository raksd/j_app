from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from app.models import StructuredCV

FONT = "Calibri"


def write_cover_letter_docx(text: str, cv: StructuredCV, out_path: Path) -> Path:
    """Builds the .docx from the (possibly user-edited) letter text. Name and contact come from the CV."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)

    head = doc.add_paragraph()
    head.paragraph_format.space_after = Pt(0)
    run = head.add_run(cv.full_name)
    run.bold = True
    run.font.size = Pt(14)
    contact = [p for p in (cv.location, cv.phone, cv.email, *cv.links) if p]
    if contact:
        doc.add_paragraph(" | ".join(contact)).paragraph_format.space_after = Pt(14)

    for block in [b for b in text.replace("\r\n", "\n").split("\n\n") if b.strip()]:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(10)
        lines = block.strip().split("\n")
        for i, line in enumerate(lines):
            para.add_run(line)
            if i < len(lines) - 1:
                para.add_run().add_break()

    doc.save(str(out_path))
    return out_path
